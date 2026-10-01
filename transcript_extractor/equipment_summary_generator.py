"""
equipment_summary_generator.py

Reads all transcript .docx files from SharePoint, scans for equipment mentions
(pumps, valves, chillers, etc.), and generates a consolidated
ChatSTP_Equipment_Master_Summary.docx uploaded back to the same Transcripts folder.
"""

import os
import io
import re
import logging
from datetime import datetime, timezone
from collections import defaultdict
from pathlib import Path

# Load .env — checks this folder first, then parent folder
try:
    from dotenv import load_dotenv
    _env = Path(__file__).resolve().parent / ".env"
    if not _env.exists():
        _env = Path(__file__).resolve().parent.parent / ".env"
    load_dotenv(dotenv_path=_env, override=True)
except ImportError:
    pass

import requests
from msal import ConfidentialClientApplication
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

logger = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# Config
# ---------------------------------------------------------------------------
TENANT_ID = os.environ["TENANT_ID"]
CLIENT_ID = os.environ["CLIENT_ID"]
CLIENT_SECRET = os.environ["CLIENT_SECRET"]
SITE_ID = os.environ["SHAREPOINT_SITE_ID"]
DRIVE_ID = os.environ["SHAREPOINT_DRIVE_ID"]
TRANSCRIPTS_ROOT = os.environ.get("TRANSCRIPTS_ROOT", "Transcripts")
SUMMARY_FILENAME = "ChatSTP_Equipment_Master_Summary.docx"

GRAPH_BASE = "https://graph.microsoft.com/v1.0"
SCOPE = ["https://graph.microsoft.com/.default"]

EQUIPMENT_PATTERNS = {
    "Pumps": [r"\b[\w\s\-]*pump[\w\s\-]*\b"],
    "Valves": [r"\b[\w\s\-]*valve[\w\s\-]*\b", r"\bMOV[\w\-]*\b", r"\bAOV[\w\-]*\b"],
    "Chillers": [r"\b[\w\s\-]*chiller[\w\s\-]*\b"],
    "Generators": [r"\b[\w\s\-]*generator[\w\s\-]*\b"],
    "Diesels": [r"\bstandby diesel[\w\s\-]*\b", r"\bemergency diesel[\w\s\-]*\b"],
    "Breakers / Transformers": [r"\b[\w\s\-]*breaker[\w\s\-]*\b", r"\b[\w\s\-]*transformer[\w\s\-]*\b"],
}


# ---------------------------------------------------------------------------
# Auth
# ---------------------------------------------------------------------------
def _get_token() -> str:
    app = ConfidentialClientApplication(
        CLIENT_ID,
        authority=f"https://login.microsoftonline.com/{TENANT_ID}",
        client_credential=CLIENT_SECRET,
    )
    result = app.acquire_token_for_client(scopes=SCOPE)
    if "access_token" not in result:
        raise RuntimeError(f"Auth failed: {result.get('error_description')}")
    return result["access_token"]


def _headers(token: str) -> dict:
    return {"Authorization": f"Bearer {token}", "Accept": "application/json"}


# ---------------------------------------------------------------------------
# SharePoint helpers
# ---------------------------------------------------------------------------
def _list_docx_files(token: str) -> list[dict]:
    files = []
    url = f"{GRAPH_BASE}/drives/{DRIVE_ID}/root:/{TRANSCRIPTS_ROOT}:/children"
    _recurse(token, url, files)
    return files


def _recurse(token: str, url: str, files: list):
    resp = requests.get(url, headers=_headers(token))
    resp.raise_for_status()
    data = resp.json()
    for item in data.get("value", []):
        if item.get("folder"):
            child_url = f"{GRAPH_BASE}/drives/{DRIVE_ID}/items/{item['id']}/children"
            _recurse(token, child_url, files)
        elif item.get("file") and item["name"].endswith(".docx") and item["name"] != SUMMARY_FILENAME:
            files.append({"id": item["id"], "name": item["name"], "webUrl": item.get("webUrl", "")})
    next_link = data.get("@odata.nextLink")
    if next_link:
        _recurse(token, next_link, files)


def _download_docx(token: str, item_id: str) -> Document:
    url = f"{GRAPH_BASE}/drives/{DRIVE_ID}/items/{item_id}/content"
    resp = requests.get(url, headers=_headers(token))
    resp.raise_for_status()
    return Document(io.BytesIO(resp.content))


def _upload_docx(token: str, doc: Document, folder_path: str):
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    upload_url = f"{GRAPH_BASE}/drives/{DRIVE_ID}/root:/{folder_path}/{SUMMARY_FILENAME}:/content"
    resp = requests.put(
        upload_url,
        headers={**_headers(token), "Content-Type": "application/octet-stream"},
        data=buffer.read(),
    )
    resp.raise_for_status()
    logger.info(f"Uploaded summary: {SUMMARY_FILENAME} to {folder_path}/")


# ---------------------------------------------------------------------------
# Parsing helpers
# ---------------------------------------------------------------------------
def _parse_meeting_name_and_date(filename: str) -> tuple[str, str]:
    name = filename.replace(".docx", "")
    parts = re.split(r"_?-?_?AutoRecord", name, maxsplit=1)
    meeting_name = parts[0].replace("_", " ").strip(" -")
    date_match = re.search(r"(\d{8})", filename)
    if date_match:
        raw = date_match.group(1)
        try:
            dt = datetime.strptime(raw, "%Y%m%d")
            date_str = dt.strftime("%B %d, %Y")
        except ValueError:
            date_str = raw
    else:
        # Try M-DD or M-D format e.g. "2-24" = February 24, 2026
        date_match2 = re.search(r"(\d{1,2})-(\d{1,2})", filename)
        if date_match2:
            try:
                m, d = int(date_match2.group(1)), int(date_match2.group(2))
                dt = datetime(2026, m, d)
                date_str = dt.strftime("%B %d, %Y")
            except ValueError:
                date_str = "Unknown Date"
        else:
            date_str = "Unknown Date"
    return meeting_name, date_str


def _extract_text(doc: Document) -> str:
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def _find_equipment(text: str) -> dict[str, list[str]]:
    results = defaultdict(set)
    lower_text = text.lower()
    for category, patterns in EQUIPMENT_PATTERNS.items():
        for pattern in patterns:
            for m in re.finditer(pattern, lower_text):
                matched_term = text[m.start():m.end()].strip()
                if matched_term:
                    results[category].add(matched_term.title())
    return {k: sorted(v) for k, v in results.items() if v}


# ---------------------------------------------------------------------------
# Document builder
# ---------------------------------------------------------------------------
def _build_summary_doc(findings: list[dict]) -> Document:
    doc = Document()

    title = doc.add_heading("ChatSTP Equipment Master Summary", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.runs[0].font.color.rgb = RGBColor(0x00, 0x3A, 0x70)

    ts = doc.add_paragraph()
    ts.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ts_run = ts.add_run(f"Auto-generated: {datetime.now(timezone.utc).strftime('%B %d, %Y at %H:%M UTC')}")
    ts_run.font.size = Pt(9)
    ts_run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
    doc.add_paragraph()

    toc = doc.add_paragraph()
    toc.add_run(
        f"This document consolidates equipment mentions across {len(findings)} meeting transcript(s). "
        "It is regenerated automatically each time new transcripts are uploaded."
    ).italic = True
    doc.add_paragraph()

    by_category = defaultdict(list)
    for entry in findings:
        for category, items in entry["equipment"].items():
            by_category[category].append({
                "meeting": entry["meeting"], "date": entry["date"],
                "url": entry["url"], "items": items,
            })

    for category in EQUIPMENT_PATTERNS.keys():
        if category not in by_category:
            continue
        doc.add_heading(category, level=1)
        for entry in sorted(by_category[category], key=lambda x: x["date"]):
            mtg_para = doc.add_paragraph(style="List Bullet")
            mtg_para.add_run(f"{entry['meeting']} — {entry['date']}").bold = True
            for item in entry["items"]:
                doc.add_paragraph(style="List Bullet 2").add_run(item)
        doc.add_paragraph()

    doc.add_heading("Full Breakdown by Meeting", level=1)
    for entry in sorted(findings, key=lambda x: x["date"]):
        doc.add_heading(f"{entry['meeting']} — {entry['date']}", level=2)
        if not entry["equipment"]:
            doc.add_paragraph("No equipment keywords detected in this transcript.")
            continue
        for category, items in entry["equipment"].items():
            p = doc.add_paragraph()
            p.add_run(f"{category}: ").bold = True
            p.add_run(", ".join(items))
        doc.add_paragraph()

    return doc


# ---------------------------------------------------------------------------
# Main entry point
# ---------------------------------------------------------------------------
async def generate_equipment_summary():
    logger.info("Starting equipment summary generation...")
    token = _get_token()
    files = _list_docx_files(token)
    logger.info(f"Found {len(files)} transcript files in SharePoint")

    findings = []
    for f in files:
        try:
            doc = _download_docx(token, f["id"])
            text = _extract_text(doc)
            equipment = _find_equipment(text)
            meeting_name, date_str = _parse_meeting_name_and_date(f["name"])
            findings.append({"meeting": meeting_name, "date": date_str, "url": f["webUrl"], "equipment": equipment})
            logger.info(f"Scanned: {f['name']} — {len(equipment)} categories found")
        except Exception as e:
            logger.warning(f"Could not process {f['name']}: {e}")

    if not findings:
        logger.warning("No transcripts found — skipping summary generation")
        return

    summary_doc = _build_summary_doc(findings)
    _upload_docx(token, summary_doc, TRANSCRIPTS_ROOT)
    logger.info("Equipment summary generation complete.")


# ---------------------------------------------------------------------------
# Standalone test runner
# ---------------------------------------------------------------------------
if __name__ == "__main__":
    import asyncio
    logging.basicConfig(level=logging.INFO)
    asyncio.run(generate_equipment_summary())
