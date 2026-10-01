import asyncio
import calendar
import hashlib
import logging
import os
import re
from datetime import datetime, timedelta, timezone
from io import BytesIO
from pathlib import Path
from typing import Any, Dict, List, Literal, Optional

import httpx
from azure.core.exceptions import ResourceExistsError
from azure.data.tables import TableClient, TableServiceClient
from azure.identity import ClientSecretCredential
from docx import Document
from docx.shared import Pt
from msgraph import GraphServiceClient

# ---------------------------------------------------------------------------
# Local .env support
# ---------------------------------------------------------------------------
try:
    from dotenv import load_dotenv
    _env = None
    _search = Path(__file__).resolve().parent
    for _ in range(4):
        candidate = _search / ".env"
        if candidate.exists():
            _env = candidate
            break
        _search = _search.parent
    if _env:
        load_dotenv(dotenv_path=_env, override=True)
        print(f"[INFO] Loaded .env from: {_env}")
    else:
        print("[WARNING] No .env file found")
except ImportError:
    pass

# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------
TENANT_ID     = os.environ.get("TENANT_ID", "")
CLIENT_ID     = os.environ.get("CLIENT_ID", "")
CLIENT_SECRET = os.environ.get("CLIENT_SECRET", "")
SINGLE_USER_ID = os.environ.get("SINGLE_USER_ID")

UserMode = Literal["single", "all", "groups", "users"]

ALLOWED_TEMPLATE_IDS: set = {
    t.strip()
    for t in os.environ.get(
        "AUTORECORD_TEMPLATE_IDS",
        "customtemplate_b1692ce2-550d-405d-b98e-19174006b2a1,"
        "customtemplate_fe60e884-8e66-4e9c-82d4-6cd8e1a08139",
    ).split(",")
    if t.strip()
}

MAX_TRANSCRIPT_AGE_DAYS  = int(os.environ.get("MAX_TRANSCRIPT_AGE_DAYS", "90"))
SHAREPOINT_SITE_ID       = os.environ.get("SHAREPOINT_SITE_ID")
SHAREPOINT_DRIVE_ID      = os.environ.get("SHAREPOINT_DRIVE_ID")
SHAREPOINT_FOLDER_PATH   = os.environ.get("SHAREPOINT_FOLDER_PATH", "Transcripts")
USE_SHAREPOINT           = os.environ.get("USE_SHAREPOINT", "false").lower() == "true"
FOLDER_GRANULARITY       = os.environ.get("FOLDER_GRANULARITY", "month").lower()
STORAGE_CONN_STR         = os.environ.get("AzureWebJobsStorage")
PROCESSED_TABLE          = "ProcessedRecordings"
NO_TRANSCRIPT_SKIP_HOURS = float(os.environ.get("NO_TRANSCRIPT_SKIP_HOURS", "4"))

MIN_RECORDING_MINUTES = float(os.environ.get("MIN_RECORDING_MINUTES", "3"))

_FORCE_REPROCESS_KEYWORDS = [
    k.strip()
    for k in os.environ.get("FORCE_REPROCESS_KEYWORDS", "").split(",")
    if k.strip()
]


def _force_norm(s: str) -> str:
    return re.sub(r"[\s_]+", "_", s or "")


_FORCE_REPROCESS_KEYWORDS_NORM = [_force_norm(k) for k in _FORCE_REPROCESS_KEYWORDS]

FETCH_ATTENDANCE_ROSTER = os.environ.get("FETCH_ATTENDANCE_ROSTER", "false").lower() == "true"

ATTENDEE_SOURCE = os.environ.get("ATTENDEE_SOURCE", "speakers").strip().lower()
if ATTENDEE_SOURCE not in ("speakers", "calendar", "attendance"):
    ATTENDEE_SOURCE = "speakers"

NORMALIZE_ATTENDEE_NAMES = os.environ.get("NORMALIZE_ATTENDEE_NAMES", "true").lower() == "true"

DIRECT_MATCH_WINDOW_HOURS = float(os.environ.get("DIRECT_MATCH_WINDOW_HOURS", "6"))
ATTENDANCE_MATCH_WINDOW_HOURS = float(os.environ.get("ATTENDANCE_MATCH_WINDOW_HOURS", "2"))
ATTENDEE_MERGE_SPEAKERS = os.environ.get("ATTENDEE_MERGE_SPEAKERS", "true").lower() == "true"
ATTENDEE_INCLUDE_PHONE = os.environ.get("ATTENDEE_INCLUDE_PHONE", "true").lower() == "true"
RESOLVE_PHONE_TO_NAME = os.environ.get("RESOLVE_PHONE_TO_NAME", "false").lower() == "true"

KNOWN_PHONE_CALLERS: Dict[str, str] = {
    k.strip(): v.strip()
    for k, v in (
        pair.split("=", 1)
        for pair in os.environ.get("KNOWN_PHONE_CALLERS", "").split(";")
        if "=" in pair
    )
}


def get_user_mode_from_env() -> tuple:
    mode = os.environ.get("RECORDING_USER_MODE", "single").lower()
    group_ids = [g.strip() for g in os.environ.get("RECORDING_GROUP_IDS", "").split(",") if g.strip()]
    if mode not in ("single", "all", "groups", "users"):
        raise ValueError(f"Unsupported RECORDING_USER_MODE: {mode}")
    return mode, SINGLE_USER_ID, group_ids


async def get_access_token(credential: ClientSecretCredential) -> Optional[str]:
    try:
        return credential.get_token("https://graph.microsoft.com/.default").token
    except Exception as exc:
        logging.error("Token acquisition failed: %s", exc)
        return None


def _is_teams_join_url(value: Any) -> bool:
    return (
        isinstance(value, str)
        and "teams.microsoft.com" in value
        and "meetup-join" in value
    )


def _deep_scan_for_url(obj: Any, path: str = "") -> Optional[str]:
    if isinstance(obj, dict):
        for k, v in obj.items():
            current_path = f"{path}.{k}" if path else k
            if _is_teams_join_url(v):
                logging.info("Deep-scan found Teams URL at '%s': %s", current_path, str(v)[:80])
                return v
            result = _deep_scan_for_url(v, current_path)
            if result:
                return result
    elif isinstance(obj, list):
        for i, item in enumerate(obj):
            result = _deep_scan_for_url(item, f"{path}[{i}]")
            if result:
                return result
    return None


def _extract_join_url_from_additional_data(additional: dict, item_name: str) -> Optional[str]:
    for key in ("joinWebUrl", "onlineMeetingUrl", "meetingJoinUrl", "meetingUrl", "joinUrl"):
        val = additional.get(key)
        if _is_teams_join_url(val):
            logging.info("Strategy 1: found '%s' at top level for '%s'", key, item_name)
            return val

    source = additional.get("source", {}) or {}
    for key in ("joinWebUrl", "onlineMeetingUrl", "meetingJoinUrl", "meetingUrl", "joinUrl"):
        val = source.get(key)
        if _is_teams_join_url(val):
            logging.info("Strategy 2: found source.%s for '%s'", key, item_name)
            return val

    media = additional.get("media", {}) or {}
    for key in ("joinWebUrl", "onlineMeetingUrl", "meetingJoinUrl", "meetingUrl", "joinUrl"):
        val = media.get(key)
        if _is_teams_join_url(val):
            logging.info("Strategy 2: found media.%s for '%s'", key, item_name)
            return val

    media_source = media.get("mediaSource", {}) or {}
    for key in ("joinWebUrl", "onlineMeetingUrl", "meetingJoinUrl", "meetingUrl", "joinUrl"):
        val = media_source.get(key)
        if _is_teams_join_url(val):
            logging.info("Strategy 2: found media.mediaSource.%s for '%s'", key, item_name)
            return val

    logging.info(
        "Strategy 3: deep-scanning additional_data for '%s'. source=%s | media_keys=%s | all_keys=%s",
        item_name, source, list(media.keys()), list(additional.keys()),
    )
    result = _deep_scan_for_url(additional)
    if result:
        logging.info("Strategy 3: deep-scan found URL for '%s': %s", item_name, result[:80])
        return result

    logging.info("No Teams join URL found in additional_data for '%s'.", item_name)
    return None


async def get_meeting_template_id(
    credential: ClientSecretCredential,
    user_id: str,
    join_web_url: str,
    organizer_id: Optional[str] = None,
) -> tuple[Optional[str], Optional[str]]:
    """Resolve a meeting's (meetingTemplateId, onlineMeetingId) from its joinWebUrl.

    onlineMeetings only returns meetings the *queried* user organized, so a
    recording sitting in user A's OneDrive that was actually organized by user B
    returns 404 when queried under A -- which previously caused the recording to
    be dropped at the template gate before the organizer-aware transcript fetch
    could ever run. We therefore try the recording owner first, then fall back to
    the organizer (same pattern as get_meeting_transcripts).

    The returned meetingTemplateId is a property of the meeting itself (uniquely
    keyed by joinWebUrl), so the AutoRecord template gate at the call site behaves
    identically regardless of which user resolved the lookup -- a non-AutoRecord
    meeting resolved via the organizer still returns its real, non-allowed
    template id and is rejected by the caller.
    """
    if not join_web_url:
        return None, None

    access_token = await get_access_token(credential)
    if not access_token:
        return None, None

    headers = {"Authorization": f"Bearer {access_token}", "Content-Type": "application/json"}

    # Recording owner first, then the organizer -- deduped, and only if different.
    query_user_ids: List[str] = [user_id]
    if organizer_id and organizer_id.lower() != user_id.lower():
        query_user_ids.append(organizer_id)

    for q_user in query_user_ids:
        url = (
            f"https://graph.microsoft.com/v1.0/users/{q_user}/onlineMeetings"
            f"?$filter=JoinWebUrl eq '{join_web_url}'"
        )
        try:
            async with httpx.AsyncClient() as client:
                resp = await client.get(url, headers=headers, timeout=30)

            if resp.status_code == 200:
                meetings = resp.json().get("value", [])
                if meetings:
                    template_id = meetings[0].get("meetingTemplateId")
                    meeting_id  = meetings[0].get("id")
                    logging.info(
                        "meetingTemplateId=%s (resolved under user=%s) for joinWebUrl %s",
                        template_id, q_user, join_web_url,
                    )
                    return template_id, meeting_id
                logging.warning(
                    "onlineMeetings returned empty (user=%s) for joinWebUrl: %s",
                    q_user, join_web_url,
                )
            else:
                logging.warning(
                    "onlineMeetings query failed (%s) (user=%s) for joinWebUrl %s: %s",
                    resp.status_code, q_user, join_web_url, resp.text[:300],
                )
        except Exception as exc:
            logging.warning("get_meeting_template_id exception (user=%s): %s", q_user, exc)

    return None, None


def _extract_join_url_from_event(event: dict) -> Optional[str]:
    online_mtg_url = event.get("onlineMeetingUrl", "")
    if online_mtg_url and "teams.microsoft.com" in online_mtg_url:
        return online_mtg_url
    join_url_obj = (event.get("onlineMeeting") or {}).get("joinUrl", "")
    if join_url_obj and "teams.microsoft.com" in join_url_obj:
        return join_url_obj
    return None


_series_master_cache: dict[str, str] = {}


def _normalize_subject(subject: str) -> str:
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9 ]", "", subject.lower())).strip()


def _extract_recording_subject(recording_name: str) -> str:
    title = recording_name
    for ext in (".mp4", ".webm", ".mkv"):
        if title.lower().endswith(ext):
            title = title[: -len(ext)]
            break
    title = re.sub(r"[\s_\-]*\d{8}_\d{6}(?:UTC)?[\s_\-]*", " ", title, flags=re.IGNORECASE)
    title = re.sub(r"[\s_\-]*\d{8}\s+\d{6}UTC[\s_\-]*", " ", title, flags=re.IGNORECASE)
    title = re.sub(r"meeting[\s_\-]*recording", "", title, flags=re.IGNORECASE)
    title = re.sub(r"autorecord", "", title, flags=re.IGNORECASE)
    return _normalize_subject(title)


def _format_attendee_name(name: str) -> str:
    if not name:
        return ""
    name = " ".join(str(name).strip().split())
    if not name or re.fullmatch(r"@\d+", name):
        return ""

    def _fix_case(s: str) -> str:
        s = s.strip()
        return s.title() if (s == s.upper() and any(c.isalpha() for c in s)) else s

    if re.search(r"\d", name):
        return _fix_case(name)

    if "," in name:
        last, _, first = name.partition(",")
        last, first = _fix_case(last), _fix_case(first)
        return f"{last}, {first}".strip().rstrip(",").strip()

    tokens = name.split(" ")
    if len(tokens) == 2:
        return f"{_fix_case(tokens[0])}, {_fix_case(tokens[1])}"
    return _fix_case(name)


async def _get_join_url_from_calendar(
    credential: ClientSecretCredential,
    user_id: str,
    item_created_dt: Optional[datetime],
    ical_uid: Optional[str] = None,
    thread_id: Optional[str] = None,
    organizer_id: Optional[str] = None,
    recording_name: Optional[str] = None,
) -> Optional[str]:
    access_token = await get_access_token(credential)
    if not access_token:
        return None

    headers = {"Authorization": f"Bearer {access_token}", "Content-Type": "application/json"}

    if ical_uid:
        url = (
            f"https://graph.microsoft.com/v1.0/users/{user_id}/events"
            f"?$filter=iCalUId eq '{ical_uid}'"
            f"&$select=subject,onlineMeetingUrl,isOnlineMeeting,onlineMeeting,start,end"
        )
        try:
            async with httpx.AsyncClient() as client:
                resp = await client.get(url, headers=headers, timeout=30)

            if resp.status_code == 200:
                events = resp.json().get("value", [])
                logging.info(
                    "iCalUid lookup returned %d event(s) for user %s (uid=%s)",
                    len(events), user_id, ical_uid[:40],
                )
                for event in events:
                    join_url = _extract_join_url_from_event(event)
                    if join_url:
                        logging.info(
                            "Calendar [iCalUid]: using join URL from event '%s'",
                            event.get("subject", "(no subject)"),
                        )
                        return join_url
                if events:
                    logging.info(
                        "Calendar [iCalUid]: event found but has no Teams join URL (uid=%s)",
                        ical_uid[:40],
                    )
                    return None
            elif resp.status_code == 403:
                logging.warning(
                    "Calendar 403 for user %s — Calendars.Read may not be propagated yet.", user_id,
                )
                return None
            else:
                logging.warning(
                    "iCalUid lookup failed (%s) for user %s: %s",
                    resp.status_code, user_id, resp.text[:200],
                )
        except Exception as exc:
            logging.warning("iCalUid calendar lookup exception: %s", exc)

    if thread_id and organizer_id:
        import urllib.parse
        context = urllib.parse.quote(
            f'{{"Tid":"{TENANT_ID}","Oid":"{organizer_id}"}}', safe=""
        )
        encoded_thread = urllib.parse.quote(thread_id, safe="")
        constructed_url = (
            f"https://teams.microsoft.com/l/meetup-join/"
            f"{encoded_thread}/0?context={context}"
        )
        logging.info("Calendar [threadId]: constructed joinWebUrl for thread %s", thread_id[:60])
        return constructed_url

    if not item_created_dt:
        return None

    if item_created_dt.tzinfo is None:
        item_created_dt = item_created_dt.replace(tzinfo=timezone.utc)

    start_bound = (item_created_dt - timedelta(hours=2)).strftime("%Y-%m-%dT%H:%M:%SZ")
    end_bound   = (item_created_dt + timedelta(hours=12)).strftime("%Y-%m-%dT%H:%M:%SZ")

    url = (
        f"https://graph.microsoft.com/v1.0/users/{user_id}/calendarView"
        f"?startDateTime={start_bound}&endDateTime={end_bound}"
        f"&$select=subject,onlineMeetingUrl,isOnlineMeeting,onlineMeeting,start,end,seriesMasterId"
    )

    try:
        async with httpx.AsyncClient() as client:
            resp = await client.get(url, headers=headers, timeout=30)

        if resp.status_code == 403:
            logging.warning(
                "Calendar 403 for user %s — Calendars.Read permission may not be fully propagated yet.",
                user_id,
            )
            return None

        if resp.status_code != 200:
            logging.warning(
                "Calendar failed (%s) for user %s: %s",
                resp.status_code, user_id, resp.text[:200],
            )
            return None

        events = resp.json().get("value", [])
        logging.info(
            "Calendar [time-window] returned %d event(s) for user %s in window %s – %s",
            len(events), user_id, start_bound, end_bound,
        )

        online_candidates = []
        for event in events:
            subject        = event.get("subject", "(no subject)")
            is_online      = event.get("isOnlineMeeting", False)
            online_mtg_url = event.get("onlineMeetingUrl", "")
            join_url_obj   = (event.get("onlineMeeting") or {}).get("joinUrl", "")
            start          = (event.get("start") or {}).get("dateTime", "")
            end            = (event.get("end") or {}).get("dateTime", "")
            series_id      = event.get("seriesMasterId", "")

            logging.info(
                "  Calendar event: subject='%s' | isOnlineMeeting=%s | "
                "onlineMeetingUrl='%s' | onlineMeeting.joinUrl='%s' | "
                "start=%s | end=%s | seriesMasterId=%s",
                subject, is_online,
                online_mtg_url[:80] if online_mtg_url else "",
                join_url_obj[:80] if join_url_obj else "",
                start, end,
                series_id[:40] if series_id else "",
            )

            join_url = _extract_join_url_from_event(event)
            if not join_url:
                continue

            norm_subj = _normalize_subject(subject)
            if series_id and norm_subj:
                if norm_subj not in _series_master_cache:
                    _series_master_cache[norm_subj] = series_id
                    logging.info(
                        "seriesMasterId cache: stored '%s' -> %s", norm_subj, series_id[:40]
                    )

            online_candidates.append({
                "event":      event,
                "join_url":   join_url,
                "series_id":  series_id,
                "start":      start,
                "subject":    subject,
            })

        if not online_candidates:
            logging.warning(
                "Calendar: no Teams join URL found in %d event(s) for user %s",
                len(events), user_id,
            )

        elif len(online_candidates) == 1:
            c = online_candidates[0]
            logging.info(
                "Calendar [time-window]: single candidate '%s', using it.", c["subject"]
            )
            return c["join_url"]

        else:
            rec_subj = _extract_recording_subject(recording_name) if recording_name else ""
            rec_series_id = _series_master_cache.get(rec_subj)

            logging.info(
                "Calendar [time-window]: %d candidates, recording_subject='%s', "
                "cached_seriesMasterId=%s",
                len(online_candidates), rec_subj,
                rec_series_id[:40] if rec_series_id else "None",
            )

            if rec_series_id:
                series_match = [
                    c for c in online_candidates if c["series_id"] == rec_series_id
                ]
                if len(series_match) == 1:
                    logging.info(
                        "Calendar [time-window]: seriesMasterId match -> '%s'",
                        series_match[0]["subject"],
                    )
                    return series_match[0]["join_url"]
                elif len(series_match) > 1:
                    logging.warning(
                        "Calendar [time-window]: %d events share seriesMasterId %s, "
                        "falling back to time-closest.",
                        len(series_match), rec_series_id[:40],
                    )
                    def _start_delta(c: dict) -> float:
                        try:
                            from dateutil.parser import parse as dtparse
                            ev_start = dtparse(c["start"]).replace(tzinfo=timezone.utc)
                            return abs((ev_start - item_created_dt).total_seconds())
                        except Exception:
                            return float("inf")
                    best = min(series_match, key=_start_delta)
                    logging.info(
                        "Calendar [time-window]: time-closest within series -> '%s'",
                        best["subject"],
                    )
                    return best["join_url"]
                else:
                    logging.warning(
                        "Calendar [time-window]: cached seriesMasterId %s not found "
                        "in any candidate — cache may be stale. Falling back to time-closest.",
                        rec_series_id[:40],
                    )

            def _start_delta_all(c: dict) -> float:
                try:
                    from dateutil.parser import parse as dtparse
                    ev_start = dtparse(c["start"]).replace(tzinfo=timezone.utc)
                    return abs((ev_start - item_created_dt).total_seconds())
                except Exception:
                    return float("inf")

            best = min(online_candidates, key=_start_delta_all)
            logging.warning(
                "Calendar [time-window]: no seriesMasterId available, "
                "using time-closest fallback -> '%s'",
                best["subject"],
            )
            return best["join_url"]

    except Exception as exc:
        logging.warning("Calendar fallback exception: %s", exc)

    return None


async def get_user_recordings(
    graph_client: GraphServiceClient,
    credential: ClientSecretCredential,
    user_id: str,
) -> tuple:
    try:
        drive = await graph_client.users.by_user_id(user_id).drive.get()
        logging.info("Drive for user %s: %s", user_id, drive.id)

        all_root_items = []
        page = await graph_client.drives.by_drive_id(drive.id).items.by_drive_item_id("root").children.get()
        while page:
            if page.value:
                all_root_items.extend(page.value)
            next_link = getattr(page, "odata_next_link", None)
            if not next_link:
                break
            page = await graph_client.drives.by_drive_id(drive.id).items.by_drive_item_id("root").children.with_url(next_link).get()

        folder_names = [item.name for item in all_root_items if getattr(item, "folder", None)]
        logging.info("Root folders for user %s: %s", user_id, folder_names)

        recordings_folder_id = None
        for item in all_root_items:
            if item.name and item.name.lower() == "recordings" and getattr(item, "folder", None):
                recordings_folder_id = item.id
                logging.info("Found Recordings folder: %s (id=%s)", item.name, item.id)
                break

        if not recordings_folder_id:
            logging.warning("No 'Recordings' folder for user %s.", user_id)
            return [], drive.id

        all_items = []
        page = await graph_client.drives.by_drive_id(drive.id).items.by_drive_item_id(recordings_folder_id).children.get()
        while page:
            if page.value:
                all_items.extend(page.value)
            next_link = getattr(page, "odata_next_link", None)
            if not next_link:
                break
            page = await graph_client.drives.by_drive_id(drive.id).items.by_drive_item_id(recordings_folder_id).children.with_url(next_link).get()

        recordings = []
        for item in all_items:
            if not (item.name and item.name.endswith((".mp4", ".webm", ".mkv"))):
                continue

            item_details = await graph_client.drives.by_drive_id(drive.id).items.by_drive_item_id(item.id).get()

            is_recording = False
            join_web_url = None
            start_dt     = None
            end_dt       = None
            ical_uid     = None
            thread_id    = None
            organizer_id = None

            if getattr(item_details, "additional_data", None):
                additional   = item_details.additional_data
                media        = additional.get("media", {}) or {}
                media_source = media.get("mediaSource", {}) or {}
                content_type = media_source.get("contentType")

                if content_type == "Recording":
                    is_recording = True
                    join_web_url = _extract_join_url_from_additional_data(additional, item.name)

                raw_start = media.get("recordingStartDateTime")
                raw_end   = media.get("recordingEndDateTime")

                source      = additional.get("source") or {}
                ical_uid    = None
                thread_id   = None
                organizer_id = None
                raw_uid     = source.get("iCalUid")
                if raw_uid:
                    ical_uid = str(raw_uid)
                raw_thread = source.get("threadId")
                if raw_thread:
                    thread_id = str(raw_thread)
                raw_org = source.get("meetingOrganizerId")
                if raw_org:
                    organizer_id = str(raw_org)

                if isinstance(raw_start, datetime):
                    start_dt = raw_start
                elif isinstance(raw_start, str):
                    try:
                        start_dt = datetime.fromisoformat(raw_start.replace("Z", "+00:00"))
                    except ValueError:
                        pass

                if isinstance(raw_end, datetime):
                    end_dt = raw_end
                elif isinstance(raw_end, str):
                    try:
                        end_dt = datetime.fromisoformat(raw_end.replace("Z", "+00:00"))
                    except ValueError:
                        pass

                logging.info(
                    "Item '%s' | contentType=%s | joinWebUrl=%s | start=%s | end=%s",
                    item.name, content_type,
                    (join_web_url[:60] + "...") if join_web_url else None,
                    raw_start, raw_end,
                )

            if not is_recording:
                continue

            if not join_web_url:
                created_dt = getattr(item_details, "created_date_time", None) or start_dt
                join_web_url = await _get_join_url_from_calendar(
                    credential, user_id, created_dt,
                    ical_uid=ical_uid,
                    thread_id=thread_id,
                    organizer_id=organizer_id,
                    recording_name=item.name,
                )

            if not join_web_url:
                continue

            template_id, online_meeting_id = await get_meeting_template_id(
                credential, user_id, join_web_url, organizer_id=organizer_id,
            )
            if template_id not in ALLOWED_TEMPLATE_IDS:
                continue

            logging.info("PASS [templateId=%s]: %s", template_id, item.name)
            recordings.append({
                "id":                item.id,
                "name":              item.name,
                "created":           item.created_date_time,
                "size":              item.size,
                "join_web_url":      join_web_url,
                "template_id":       template_id,
                "online_meeting_id": online_meeting_id,
                "start_dt":          start_dt,
                "end_dt":            end_dt,
                "ical_uid":          ical_uid,
                "thread_id":         thread_id,
                "organizer_id":      organizer_id,
            })

        logging.info("User %s: %d template-matched recording(s) found", user_id, len(recordings))
        return recordings, drive.id

    except Exception as exc:
        logging.exception("Error accessing recordings for user %s: %s", user_id, exc)
        return [], None


async def get_all_transcripts_for_user(credential: ClientSecretCredential, user_id: str) -> List[dict]:
    try:
        access_token = await get_access_token(credential)
        if not access_token:
            return []
        headers = {"Authorization": f"Bearer {access_token}", "Content-Type": "application/json"}
        all_transcripts = []
        url = (
            f"https://graph.microsoft.com/beta/users/{user_id}"
            f"/onlineMeetings/getAllTranscripts(meetingOrganizerUserId='{user_id}')"
        )
        async with httpx.AsyncClient() as client:
            while url:
                resp = await client.get(url, headers=headers, timeout=30.0)
                if resp.status_code == 403:
                    logging.warning("getAllTranscripts 403 for user %s", user_id)
                    return []
                if resp.status_code != 200:
                    logging.error("getAllTranscripts failed (%s): %s", resp.status_code, resp.text[:300])
                    return []
                data = resp.json()
                all_transcripts.extend(data.get("value", []))
                url = data.get("@odata.nextLink")
        logging.info("User %s: %d transcript(s)", user_id, len(all_transcripts))
        return all_transcripts
    except Exception as exc:
        logging.exception("Error fetching transcripts for user %s: %s", user_id, exc)
        return []


async def get_meeting_transcripts(
    credential: ClientSecretCredential,
    organizer_user_id: str,
    online_meeting_id: str,
) -> List[dict]:
    if not (organizer_user_id and online_meeting_id):
        return []
    access_token = await get_access_token(credential)
    if not access_token:
        return []
    headers = {"Authorization": f"Bearer {access_token}", "Content-Type": "application/json"}
    for api in ("v1.0", "beta"):
        url = (
            f"https://graph.microsoft.com/{api}/users/{organizer_user_id}"
            f"/onlineMeetings/{online_meeting_id}/transcripts"
        )
        transcripts: List[dict] = []
        ok = True
        try:
            async with httpx.AsyncClient() as client:
                while url:
                    resp = await client.get(url, headers=headers, timeout=30.0)
                    if resp.status_code != 200:
                        logging.info(
                            "Direct transcripts (%s) unavailable (%s) for meeting %s: %s",
                            api, resp.status_code, online_meeting_id[:40], resp.text[:160],
                        )
                        ok = False
                        break
                    data = resp.json()
                    transcripts.extend(data.get("value", []))
                    url = data.get("@odata.nextLink")
        except Exception as exc:
            logging.warning("Direct transcripts (%s) exception for meeting %s: %s", api, online_meeting_id[:40], exc)
            ok = False
        if ok:
            logging.info("Direct transcripts (%s): %d for meeting %s", api, len(transcripts), online_meeting_id[:40])
            return transcripts
    return []


def _transcript_ref_dt(t: dict) -> Optional[datetime]:
    return _parse_iso_dt(t.get("endDateTime", "") or t.get("createdDateTime", ""))


def _pick_meeting_transcript(
    transcripts: List[dict],
    rec_start_dt: Optional[datetime],
    rec_end_dt: Optional[datetime],
    used_transcript_ids: set,
) -> tuple:
    ref = rec_end_dt or rec_start_dt
    if ref is None or not transcripts:
        return None, None
    if ref.tzinfo is not None:
        ref = ref.astimezone(timezone.utc).replace(tzinfo=None)
    best, best_delta = None, float("inf")
    for t in transcripts:
        tid = t.get("id", "")
        if not tid or tid in used_transcript_ids:
            continue
        t_dt = _transcript_ref_dt(t)
        if t_dt is None:
            continue
        delta = abs((ref - t_dt).total_seconds())
        if delta < best_delta:
            best, best_delta = t, delta
    if best is not None and best_delta <= DIRECT_MATCH_WINDOW_HOURS * 3600:
        return best, best_delta / 60.0
    return None, None


def _reports_in_window(
    reports: List[dict],
    rec_start_dt: Optional[datetime],
    rec_end_dt: Optional[datetime],
    window_hours: float,
) -> List[dict]:
    if not (rec_start_dt and rec_end_dt):
        return []
    lo = rec_start_dt.astimezone(timezone.utc).replace(tzinfo=None) if rec_start_dt.tzinfo else rec_start_dt
    hi = rec_end_dt.astimezone(timezone.utc).replace(tzinfo=None) if rec_end_dt.tzinfo else rec_end_dt
    lo -= timedelta(hours=window_hours)
    hi += timedelta(hours=window_hours)
    matched = []
    for r in reports:
        r_start = _parse_iso_dt(r.get("meetingStartDateTime", "") or "")
        r_end   = _parse_iso_dt(r.get("meetingEndDateTime", "") or "") or r_start
        if r_start is None:
            continue
        if r_start <= hi and r_end >= lo:
            matched.append(r)
    return matched


def _normalize_phone_digits(s: str) -> str:
    digits = "".join(c for c in (s or "") if c.isdigit())
    return digits[-10:] if len(digits) >= 10 else digits


_DIRECTORY_PHONE_INDEX: Optional[Dict[str, str]] = None


async def _build_directory_phone_index(
    credential: ClientSecretCredential,
) -> Dict[str, str]:
    global _DIRECTORY_PHONE_INDEX
    if _DIRECTORY_PHONE_INDEX is not None:
        return _DIRECTORY_PHONE_INDEX

    index: Dict[str, str] = {}
    access_token = await get_access_token(credential)
    if not access_token:
        _DIRECTORY_PHONE_INDEX = index
        return index

    headers = {
        "Authorization": f"Bearer {access_token}",
        "Content-Type": "application/json",
        "ConsistencyLevel": "eventual",
    }
    url = (
        "https://graph.microsoft.com/v1.0/users"
        "?$filter=accountEnabled eq true and "
        "(mobilePhone ne null or businessPhones/$count ne 0)"
        "&$select=displayName,mobilePhone,businessPhones"
        "&$count=true&$top=999"
    )

    pages = 0
    try:
        async with httpx.AsyncClient() as client:
            while url and pages < 20:
                resp = await client.get(url, headers=headers, timeout=30.0)
                if resp.status_code != 200:
                    logging.info(
                        "Directory phone index fetch failed (%s): %s",
                        resp.status_code, resp.text[:200],
                    )
                    break
                body = resp.json()
                for u in body.get("value", []):
                    name = (u.get("displayName") or "").strip()
                    if not name:
                        continue
                    for raw in [u.get("mobilePhone")] + (u.get("businessPhones") or []):
                        key = _normalize_phone_digits(raw or "")
                        if key:
                            index.setdefault(key, name)
                url = body.get("@odata.nextLink")
                pages += 1
        logging.info("Directory phone index built: %d phone(s) mapped", len(index))
    except Exception as exc:
        logging.warning("Directory phone index exception: %s", exc)

    _DIRECTORY_PHONE_INDEX = index
    return index


async def resolve_phone_to_name(
    credential: ClientSecretCredential,
    phone_num: str,
) -> Optional[str]:
    if not phone_num:
        return None

    target = _normalize_phone_digits(phone_num)
    for k, v in KNOWN_PHONE_CALLERS.items():
        if _normalize_phone_digits(k) == target and target:
            return v
    if phone_num in KNOWN_PHONE_CALLERS:
        return KNOWN_PHONE_CALLERS[phone_num]

    if not RESOLVE_PHONE_TO_NAME:
        return None

    if not target:
        return None

    index = await _build_directory_phone_index(credential)
    name = index.get(target)
    if name:
        logging.info("Resolved phone %s -> %s via directory lookup", phone_num, name)
        return name

    logging.info("Phone %s not found in directory index (no name match)", phone_num)
    return None


_PHONE_CHARS = re.compile(r"^\+?[\d\*\-\.\(\) ]+$")


def _is_phone_entry(name: str) -> bool:
    s = (name or "").strip()
    if s.lower().startswith("phone caller"):
        return True
    digitsish = sum(c.isdigit() or c == "*" for c in s)
    return bool(s) and digitsish >= 5 and bool(_PHONE_CHARS.fullmatch(s))


async def get_meeting_attendees(
    credential: ClientSecretCredential,
    organizer_user_id: str,
    meeting_id: str,
    rec_start_dt: Optional[datetime] = None,
    rec_end_dt: Optional[datetime] = None,
) -> List[str]:
    if not (organizer_user_id and meeting_id):
        return []

    access_token = await get_access_token(credential)
    if not access_token:
        return []

    headers = {"Authorization": f"Bearer {access_token}", "Content-Type": "application/json"}

    try:
        async with httpx.AsyncClient() as client:
            reports: List[dict] = []
            for api in ("v1.0", "beta"):
                base = (
                    f"https://graph.microsoft.com/{api}/users/{organizer_user_id}"
                    f"/onlineMeetings/{meeting_id}/attendanceReports"
                )
                url, ok, page_reports = base, True, []
                while url:
                    resp = await client.get(url, headers=headers, timeout=30.0)
                    if resp.status_code != 200:
                        logging.info(
                            "attendanceReports (%s) unavailable (%s) for meeting %s: %s",
                            api, resp.status_code, meeting_id[:40], resp.text[:160],
                        )
                        ok = False
                        break
                    data = resp.json()
                    page_reports.extend(data.get("value", []))
                    url = data.get("@odata.nextLink")
                if ok:
                    reports = page_reports
                    report_base = base
                    break
            else:
                return []

            if not reports:
                logging.info("No attendance reports for meeting %s", meeting_id[:40])
                return []

            in_window = _reports_in_window(
                reports, rec_start_dt, rec_end_dt, ATTENDANCE_MATCH_WINDOW_HOURS
            )
            if not in_window:
                logging.info(
                    "No attendance report overlaps the recording window for meeting %s "
                    "(%d report(s) on the series) — falling back to speakers",
                    meeting_id[:40], len(reports),
                )
                return []

            records: List[dict] = []
            for report in in_window:
                report_id = report.get("id", "")
                if not report_id:
                    continue
                detail_url = f"{report_base}/{report_id}?$expand=attendanceRecords"
                resp2 = await client.get(detail_url, headers=headers, timeout=30.0)
                if resp2.status_code != 200:
                    logging.info(
                        "attendanceRecords fetch failed (%s) for report %s",
                        resp2.status_code, report_id[:40],
                    )
                    continue
                records.extend(resp2.json().get("attendanceRecords", []) or [])

        logging.info(
            "Meeting %s: %d/%d report(s) in recording window, %d raw attendance record(s) merged",
            meeting_id[:40], len(in_window), len(reports), len(records),
        )

        totals: Dict[str, float] = {}
        display_for: Dict[str, str] = {}
        for rec in records:
            ident = rec.get("identity") or {}
            display = (ident.get("displayName") or rec.get("emailAddress") or "").strip()

            # FIX: Microsoft sometimes puts raw phone digits in displayName for PSTN callers.
            # If displayName looks like a phone number, clear it so the phone branch fires.
            if display and _is_phone_entry(display):
                display = ""

            if not display:
                phone_obj = ident.get("phone") or {}
                phone_num = (phone_obj.get("id") or phone_obj.get("phoneNumber") or "").strip()
                if not phone_num:
                    phone_num = (rec.get("phoneNumber") or "").strip()
                if phone_num:
                    resolved_name = await resolve_phone_to_name(credential, phone_num)
                    if resolved_name:
                        display = resolved_name
                    else:
                        display_num = phone_num if phone_num.startswith("+") else f"+{phone_num}"
                        display = f"Phone caller ({display_num})"
            if not display:
                continue
            key = display.lower()
            totals[key] = totals.get(key, 0) + (rec.get("totalAttendanceInSeconds", 0) or 0)
            display_for.setdefault(key, display)

        names = [display_for[k] for k in sorted(totals, key=lambda k: totals[k], reverse=True)]
        logging.info(
            "Attendance roster for meeting %s: %d attendee(s)", meeting_id[:40], len(names)
        )
        return names

    except Exception as exc:
        logging.warning("get_meeting_attendees exception for meeting %s: %s", meeting_id[:40], exc)
        return []


def _extract_event_attendee_names(event: dict) -> List[str]:
    attendees = event.get("attendees", []) or []
    names: List[str] = []
    seen: set = set()
    for a in attendees:
        email = a.get("emailAddress", {}) or {}
        name = (email.get("name") or "").strip()
        if not name:
            continue
        key = name.lower()
        if key not in seen:
            seen.add(key)
            names.append(name)
    logging.info("Calendar event attendees: %d name(s) extracted", len(names))
    return names


async def get_calendar_event_attendees(
    credential: ClientSecretCredential,
    user_id: str,
    ical_uid: Optional[str] = None,
    start_dt: Optional[datetime] = None,
) -> List[str]:
    access_token = await get_access_token(credential)
    if not access_token:
        return []
    headers = {"Authorization": f"Bearer {access_token}", "Content-Type": "application/json"}

    if ical_uid:
        url = (
            f"https://graph.microsoft.com/v1.0/users/{user_id}/events"
            f"?$filter=iCalUId eq '{ical_uid}'"
            f"&$select=attendees"
        )
        try:
            async with httpx.AsyncClient() as client:
                resp = await client.get(url, headers=headers, timeout=30)
            if resp.status_code == 200:
                events = resp.json().get("value", [])
                if events:
                    return _extract_event_attendee_names(events[0])
        except Exception as exc:
            logging.warning("Calendar attendee iCalUid lookup exception: %s", exc)

    if start_dt:
        if start_dt.tzinfo is None:
            start_dt = start_dt.replace(tzinfo=timezone.utc)
        start_bound = (start_dt - timedelta(hours=2)).strftime("%Y-%m-%dT%H:%M:%SZ")
        end_bound   = (start_dt + timedelta(hours=12)).strftime("%Y-%m-%dT%H:%M:%SZ")
        url = (
            f"https://graph.microsoft.com/v1.0/users/{user_id}/calendarView"
            f"?startDateTime={start_bound}&endDateTime={end_bound}"
            f"&$select=attendees,start"
        )
        try:
            async with httpx.AsyncClient() as client:
                resp = await client.get(url, headers=headers, timeout=30)
            if resp.status_code == 200:
                events = resp.json().get("value", [])
                best = None
                best_diff = float("inf")
                for ev in events:
                    ev_start = _parse_iso_dt(ev.get("start", {}).get("dateTime", ""))
                    if ev_start:
                        if start_dt.tzinfo is not None:
                            ev_start = ev_start.replace(tzinfo=timezone.utc)
                        diff = abs((start_dt - ev_start).total_seconds())
                        if diff < best_diff:
                            best_diff = diff
                            best = ev
                if best:
                    return _extract_event_attendee_names(best)
        except Exception as exc:
            logging.warning("Calendar attendee time-window lookup exception: %s", exc)

    return []


def _parse_recording_dt(recording_name: str) -> Optional[datetime]:
    m = re.search(r"(\d{8}_\d{6})(?:UTC)?", recording_name, flags=re.IGNORECASE)
    if m:
        try:
            return datetime.strptime(m.group(1), "%Y%m%d_%H%M%S")
        except ValueError:
            pass
    m = re.search(r"(\d{8})\s+(\d{6})UTC", recording_name, flags=re.IGNORECASE)
    if m:
        try:
            return datetime.strptime(m.group(1) + "_" + m.group(2), "%Y%m%d_%H%M%S")
        except ValueError:
            pass
    return None


_SP_ILLEGAL_CHARS = re.compile(r'[\\/:*?"<>|#%&{}~]+')


def _extract_short_meeting_name(recording_name: str) -> str:
    title = recording_name
    for ext in (".mp4", ".webm", ".mkv"):
        if title.lower().endswith(ext):
            title = title[: -len(ext)]
            break

    title = re.sub(r"[\s_\-]*\d{8}_\d{6}(?:UTC)?[\s_\-]*", "_", title, flags=re.IGNORECASE)
    title = re.sub(r"[\s_\-]*\d{8}\s+\d{6}UTC[\s_\-]*", "_", title, flags=re.IGNORECASE)
    title = re.sub(r"[\s_\-]*meeting[\s_\-]*recording[\s_\-]*", "", title, flags=re.IGNORECASE)
    title = re.sub(r"[\s_\-]+autorecord\b", "", title, flags=re.IGNORECASE)
    title = re.sub(r"_+", "_", title).strip("_-. ")

    paren_match = re.search(r"\(([A-Z]{2,})\)", title)
    if paren_match:
        return paren_match.group(1)

    clean = title.replace(" ", "_")
    clean = _SP_ILLEGAL_CHARS.sub("_", clean)
    clean = re.sub(r"_+", "_", clean).strip("_")
    return clean if clean else "Untitled"


def build_output_basename(recording_name: str, fallback_dt: datetime) -> str:
    dt = _parse_recording_dt(recording_name) or fallback_dt
    if hasattr(dt, "tzinfo") and dt.tzinfo is not None:
        dt = dt.astimezone(timezone.utc).replace(tzinfo=None)
    date_part = dt.strftime("%Y-%m-%d")
    time_part = dt.strftime("%H%M%S")

    meeting_name = _extract_short_meeting_name(recording_name)
    safe = _SP_ILLEGAL_CHARS.sub("_", meeting_name)
    safe = re.sub(r"_+", "_", safe).strip("_")
    if not safe:
        safe = "Untitled"

    return f"{date_part}_{safe}_{time_part}"


def _parse_iso_dt(s: str) -> Optional[datetime]:
    if not s:
        return None
    try:
        return datetime.strptime(s[:19].replace("T", " "), "%Y-%m-%d %H:%M:%S")
    except ValueError:
        return None


def match_transcript_to_recording(
    transcript: dict,
    recording_name: str,
    rec_start_dt: Optional[datetime] = None,
) -> bool:
    if rec_start_dt is not None:
        rec_dt = rec_start_dt.replace(tzinfo=None) if rec_start_dt.tzinfo else rec_start_dt
    else:
        rec_dt = _parse_recording_dt(recording_name)
    if not rec_dt:
        return False

    meet_start = _parse_iso_dt(transcript.get("meetingStartDateTime", ""))
    if meet_start is not None:
        if abs((rec_dt - meet_start).total_seconds()) <= 5400:
            return True
        return False

    trans_dt = _parse_iso_dt(transcript.get("createdDateTime", ""))
    if trans_dt is not None:
        if abs((rec_dt - trans_dt).total_seconds()) <= 900:
            return True

    return False


def find_closest_transcript(
    all_transcripts: List[dict],
    recording_name: str,
    rec_start_dt: Optional[datetime] = None,
) -> Optional[dict]:
    if rec_start_dt is not None:
        rec_dt = rec_start_dt.replace(tzinfo=None) if rec_start_dt.tzinfo else rec_start_dt
    else:
        rec_dt = _parse_recording_dt(recording_name)
    if not rec_dt or not all_transcripts:
        return None
    best = None
    best_diff = float("inf")
    for t in all_transcripts:
        ref_str = t.get("meetingStartDateTime", "") or t.get("createdDateTime", "")
        trans_dt = _parse_iso_dt(ref_str)
        if trans_dt is None:
            continue
        diff = abs((rec_dt - trans_dt).total_seconds())
        if diff < best_diff:
            best_diff = diff
            best = t
    return best


def _transcript_time_delta(
    transcript: dict,
    recording_name: str,
    rec_start_dt: Optional[datetime] = None,
) -> float:
    if rec_start_dt is not None:
        rec_dt = rec_start_dt.replace(tzinfo=None) if rec_start_dt.tzinfo else rec_start_dt
    else:
        rec_dt = _parse_recording_dt(recording_name)
    if not rec_dt:
        return float("inf")
    ref_str  = transcript.get("meetingStartDateTime", "") or transcript.get("createdDateTime", "")
    trans_dt = _parse_iso_dt(ref_str)
    if trans_dt is None:
        return float("inf")
    return abs((rec_dt - trans_dt).total_seconds())


async def download_transcript_content(
    credential: ClientSecretCredential,
    user_id: str,
    meeting_id: str,
    transcript_id: str,
) -> Optional[str]:
    access_token = await get_access_token(credential)
    if not access_token:
        return None
    headers = {"Authorization": f"Bearer {access_token}", "Accept": "text/vtt"}
    attempts = []
    if meeting_id:
        attempts.append(("v1.0/meetingId", f"https://graph.microsoft.com/v1.0/users/{user_id}/onlineMeetings/{meeting_id}/transcripts/{transcript_id}/content"))
        attempts.append(("beta/meetingId", f"https://graph.microsoft.com/beta/users/{user_id}/onlineMeetings/{meeting_id}/transcripts/{transcript_id}/content"))
    attempts.append(("beta/getAllTranscripts", f"https://graph.microsoft.com/beta/users/{user_id}/onlineMeetings/getAllTranscripts/{transcript_id}/content"))
    attempts.append(("beta/transcriptId-only", f"https://graph.microsoft.com/beta/users/{user_id}/onlineMeetings/transcripts/{transcript_id}/content"))
    try:
        async with httpx.AsyncClient() as client:
            for label, url in attempts:
                try:
                    resp = await client.get(url, headers=headers, timeout=30.0)
                    if resp.status_code == 200:
                        logging.info("Transcript downloaded via %s", label)
                        return resp.text
                    logging.warning("Download %s failed (%s)", label, resp.status_code)
                except Exception as exc:
                    logging.warning("Download %s exception: %s", label, exc)
    except Exception as exc:
        logging.error("Transcript download error: %s", exc)
    logging.error("All download approaches failed for transcript %s", transcript_id)
    return None


def parse_vtt_to_lines(vtt_content: str) -> List[dict]:
    entries = []
    current_speaker = None
    current_text = []

    tagged     = re.compile(r"<v ([^>]+)>(.*?)</v>", re.DOTALL)
    plain_label = re.compile(r"^([A-Za-z][A-Za-z0-9\s'\-\.]{1,59}):\s*(.*)$")

    def flush():
        if current_speaker is not None and current_text:
            entries.append({
                "speaker": current_speaker,
                "text": " ".join(t for t in current_text if t)
            })

    for raw_line in vtt_content.split("\n"):
        line = raw_line.strip()

        if (not line
                or line.startswith("WEBVTT")
                or line.startswith("NOTE")
                or "-->" in line
                or re.match(r"^\d+$", line)):
            continue

        m_tag = tagged.match(line)
        if m_tag:
            speaker = m_tag.group(1).strip()
            if re.match(r'^@\d+$', speaker):
                if current_speaker is not None:
                    text = re.sub(r"<[^>]+>", "", m_tag.group(2)).strip()
                    if text:
                        current_text.append(text)
                continue
            text    = re.sub(r"<[^>]+>", "", m_tag.group(2)).strip()
            if speaker == current_speaker:
                current_text.append(text)
            else:
                flush()
                current_speaker = speaker
                current_text    = [text]
            continue

        m_plain = plain_label.match(line)
        if m_plain:
            candidate = m_plain.group(1).strip()
            inline    = m_plain.group(2).strip()
            if not re.fullmatch(r"\d+", candidate):
                flush()
                current_speaker = candidate
                current_text    = [inline] if inline else []
                continue

        if current_speaker is not None:
            cleaned = re.sub(r"<[^>]+>", "", line).strip()
            if cleaned:
                current_text.append(cleaned)

    flush()
    return entries


def _attendee_dedup_key(name: str) -> str:
    s = (name or "").strip().lower()
    if s.startswith("phone caller"):
        s = s[len("phone caller"):]
    s = s.strip("() ")
    if _is_phone_entry(s) or _is_phone_entry(name):
        if "*" in s:
            tail = s.rsplit("*", 1)[-1]
            tail_digits = "".join(c for c in tail if c.isdigit())
        else:
            digits_only = "".join(c for c in s if c.isdigit())
            tail_digits = digits_only[-2:] if len(digits_only) >= 2 else digits_only
        return f"phone:{tail_digits}" if tail_digits else "phone:" + s
    return "".join(c for c in s if c.isalnum())


def _format_roster_name(name: str) -> str:
    if not name:
        return ""
    name = " ".join(str(name).strip().split())
    if not name or re.fullmatch(r"@\d+", name):
        return ""

    def _fix_case(s: str) -> str:
        s = s.strip()
        return s.title() if (s == s.upper() and any(c.isalpha() for c in s)) else s

    if "," in name:
        parts = [p.strip() for p in name.split(",")]
        return ", ".join(_fix_case(p) for p in parts if p)
    return _fix_case(name)


def _resolve_final_attendees(
    roster_attendees: Optional[List[str]],
    transcript_lines: List[dict],
) -> List[str]:
    roster_attendees = roster_attendees or []
    names: List[str] = []
    phones: List[str] = []
    seen: set = set()

    def _add(raw: str, formatter) -> None:
        raw = (raw or "").strip()
        if not raw:
            return
        if _is_phone_entry(raw):
            key = _attendee_dedup_key(raw)
            if not key or key in seen:
                return
            seen.add(key)
            if ATTENDEE_INCLUDE_PHONE:
                if raw.lower().startswith("phone caller"):
                    label = raw
                else:
                    # FIX: always normalize to +number format
                    num_display = raw if raw.startswith("+") else f"+{raw}"
                    label = f"Phone caller ({num_display})"
                phones.append(label)
            return
        name = formatter(raw)
        if not name:
            return
        key = _attendee_dedup_key(name)
        if not key or key in seen:
            return
        seen.add(key)
        names.append(name)

    def _speaker_formatter(raw: str) -> str:
        if NORMALIZE_ATTENDEE_NAMES:
            return _format_attendee_name(raw)
        name = " ".join((raw or "").split())
        return "" if re.fullmatch(r"@\d+", name) else name

    for raw in roster_attendees:
        _add(raw, _format_roster_name)

    if ATTENDEE_MERGE_SPEAKERS or not roster_attendees:
        for e in transcript_lines:
            _add(e.get("speaker", ""), _speaker_formatter)

    return names + phones


def build_docx_bytes(
    meeting_title: str,
    meeting_date: datetime,
    transcript_lines: List[dict],
    vtt_content: str = "",
    roster_attendees: Optional[List[str]] = None,
) -> bytes:
    doc     = Document()
    heading = doc.add_heading(meeting_title, level=1)
    heading.runs[0].font.size = Pt(16)
    date_para = doc.add_paragraph()
    date_para.add_run("Date: ").bold = True
    date_para.add_run(meeting_date.strftime("%B %d, %Y"))

    attendees = _resolve_final_attendees(roster_attendees, transcript_lines)

    if attendees:
        att = doc.add_paragraph()
        att.add_run("Attendees: ").bold = True
        att.add_run(", ".join(attendees))
    doc.add_paragraph("\u2500" * 60)

    formatted_heading = doc.add_heading("Transcript", level=2)
    formatted_heading.runs[0].font.size = Pt(13)
    for entry in transcript_lines:
        para = doc.add_paragraph()
        para.add_run(f"{entry['speaker']}: ").bold = True
        para.add_run(entry["text"])

    if vtt_content:
        doc.add_paragraph("\u2500" * 60)
        raw_heading = doc.add_heading("Raw Transcript (VTT)", level=2)
        raw_heading.runs[0].font.size = Pt(13)
        doc.add_paragraph(vtt_content)

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def get_folder_path(meeting_date: datetime) -> str:
    base = f"{SHAREPOINT_FOLDER_PATH}/{meeting_date.year}/{calendar.month_name[meeting_date.month]}"
    if FOLDER_GRANULARITY == "day":
        base += f"/{str(meeting_date.day).zfill(2)}"
    return base


async def ensure_folder_exists(credential: ClientSecretCredential, folder_path: str) -> bool:
    if not (SHAREPOINT_SITE_ID and SHAREPOINT_DRIVE_ID):
        return False
    access_token = await get_access_token(credential)
    if not access_token:
        return False
    headers = {"Authorization": f"Bearer {access_token}", "Content-Type": "application/json"}
    parts = folder_path.split("/")
    current_path = ""
    async with httpx.AsyncClient() as client:
        for part in parts:
            parent       = current_path if current_path else "root"
            current_path = f"{current_path}/{part}" if current_path else part
            check_url    = f"https://graph.microsoft.com/v1.0/sites/{SHAREPOINT_SITE_ID}/drives/{SHAREPOINT_DRIVE_ID}/root:/{current_path}"
            resp = await client.get(check_url, headers=headers, timeout=30)
            if resp.status_code == 404:
                create_url = (
                    f"https://graph.microsoft.com/v1.0/sites/{SHAREPOINT_SITE_ID}/drives/{SHAREPOINT_DRIVE_ID}/root/children"
                    if parent == "root" else
                    f"https://graph.microsoft.com/v1.0/sites/{SHAREPOINT_SITE_ID}/drives/{SHAREPOINT_DRIVE_ID}/root:/{parent}:/children"
                )
                cr = await client.post(create_url, headers=headers, json={"name": part, "folder": {}, "@microsoft.graph.conflictBehavior": "fail"}, timeout=30)
                if cr.status_code not in (200, 201, 409):
                    logging.error("Failed to create folder %s: %s", part, cr.text[:200])
                    return False
                logging.info("Created SharePoint folder: %s", current_path)
    return True


async def upload_to_sharepoint(credential: ClientSecretCredential, file_name: str, content: bytes, meeting_date: datetime) -> Optional[str]:
    if not (SHAREPOINT_SITE_ID and SHAREPOINT_DRIVE_ID):
        logging.warning("SharePoint not configured")
        return None
    folder_path = get_folder_path(meeting_date)
    if not await ensure_folder_exists(credential, folder_path):
        return None
    access_token = await get_access_token(credential)
    if not access_token:
        return None
    content_type = (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        if file_name.endswith(".docx") else "text/plain"
    )
    raw = content if isinstance(content, bytes) else content.encode("utf-8")
    headers = {"Authorization": f"Bearer {access_token}", "Content-Type": content_type}
    url = (
        f"https://graph.microsoft.com/v1.0/sites/{SHAREPOINT_SITE_ID}"
        f"/drives/{SHAREPOINT_DRIVE_ID}/root:/{folder_path}/{file_name}:/content"
        f"?@microsoft.graph.conflictBehavior=replace"
    )
    async with httpx.AsyncClient() as client:
        resp = await client.put(url, headers=headers, content=raw, timeout=60)
        if resp.status_code in (200, 201):
            sp_path = f"{folder_path}/{file_name}"
            logging.info("Uploaded: %s", sp_path)
            return sp_path
        logging.error("Upload failed: %s %s", resp.status_code, resp.text[:200])
    return None


def _get_table_client() -> Optional[TableClient]:
    if not STORAGE_CONN_STR:
        return None
    try:
        svc = TableServiceClient.from_connection_string(STORAGE_CONN_STR)
        try:
            svc.create_table(PROCESSED_TABLE)
        except Exception:
            pass
        return TableClient.from_connection_string(STORAGE_CONN_STR, table_name=PROCESSED_TABLE)
    except Exception as exc:
        logging.error("Failed to get TableClient: %s", exc)
        return None


def is_already_processed(user_id: str, recording_id: str) -> bool:
    if not STORAGE_CONN_STR:
        return False
    try:
        tc = TableClient.from_connection_string(STORAGE_CONN_STR, table_name=PROCESSED_TABLE)
        tc.get_entity(partition_key=user_id, row_key=recording_id)
        return True
    except Exception:
        return False


def mark_as_processed(user_id: str, recording_id: str, name: str, sp_path: Optional[str]) -> bool:
    tc = _get_table_client()
    if not tc:
        return True

    entity = {
        "PartitionKey":   user_id,
        "RowKey":         recording_id,
        "ProcessedAt":    datetime.now(timezone.utc).isoformat(),
        "RecordingName":  name,
        "SharePointPath": sp_path or "",
    }
    try:
        tc.create_entity(entity)
        logging.info("Marked as processed: %s / %s", user_id, recording_id)
        return True
    except ResourceExistsError:
        logging.info(
            "Duplicate detected (concurrent run) — skipping: %s / %s", user_id, recording_id
        )
        return False
    except Exception as exc:
        logging.error("Failed to mark processed: %s", exc)
        return True


async def process_recordings(
    graph_client: GraphServiceClient,
    credential: ClientSecretCredential,
    user_id: str,
    recordings: List[dict],
    used_content_hashes: set,
) -> dict:
    results: dict = {"successful": [], "failed": []}

    meeting_transcripts_cache: Dict[str, List[dict]] = {}
    bulk_transcripts: Optional[List[dict]] = None
    used_transcript_ids = set()

    def _rec_sort_key(r: dict):
        dt = r.get("start_dt") or _parse_recording_dt(r.get("name", ""))
        if isinstance(dt, datetime) and dt.tzinfo is not None:
            dt = dt.astimezone(timezone.utc).replace(tzinfo=None)
        return dt or datetime.min

    recordings = sorted(recordings, key=_rec_sort_key)

    for rec in recordings:
        rec_id   = rec.get("id", "")
        rec_name = rec["name"]
        logging.info("Processing '%s'", rec_name)

        start_dt_check = rec.get("start_dt")
        end_dt_check   = rec.get("end_dt")
        if start_dt_check and end_dt_check:
            duration_mins = (end_dt_check - start_dt_check).total_seconds() / 60
            if duration_mins < MIN_RECORDING_MINUTES:
                logging.info(
                    "SKIP (too short, %.1f min < %.1f min threshold): %s",
                    duration_mins, MIN_RECORDING_MINUTES, rec_name,
                )
                mark_as_processed(user_id, rec_id, rec_name, "too-short")
                continue

        _rec_name_norm = _force_norm(rec_name)
        force = any(kw in _rec_name_norm for kw in _FORCE_REPROCESS_KEYWORDS_NORM)
        if not force and is_already_processed(user_id, rec_id):
            logging.info("SKIP (already processed): %s", rec_name)
            continue
        if force:
            logging.info("FORCE REPROCESS (keyword match): %s", rec_name)

        created = rec.get("created", "")
        if isinstance(created, datetime):
            meeting_date = created
        elif isinstance(created, str) and created:
            try:
                meeting_date = datetime.strptime(created[:19], "%Y-%m-%dT%H:%M:%S")
            except (ValueError, TypeError):
                meeting_date = datetime.now(timezone.utc)
        else:
            meeting_date = datetime.now(timezone.utc)

        cutoff = datetime.now(timezone.utc) - timedelta(days=MAX_TRANSCRIPT_AGE_DAYS)
        if hasattr(meeting_date, "tzinfo") and meeting_date.tzinfo is None:
            meeting_date = meeting_date.replace(tzinfo=timezone.utc)
        if meeting_date < cutoff:
            mark_as_processed(user_id, rec_id, rec_name, "age-skip")
            continue

        matched = None
        match_delta_min = None
        rec_start_dt = rec.get("start_dt")
        rec_end_dt   = rec.get("end_dt")
        online_meeting_id = rec.get("online_meeting_id", "") or ""
        rec_organizer_id  = rec.get("organizer_id", "") or ""

        meeting_transcripts: List[dict] = []
        if online_meeting_id:
            if online_meeting_id not in meeting_transcripts_cache:
                ts = await get_meeting_transcripts(credential, user_id, online_meeting_id)
                if not ts and rec_organizer_id and rec_organizer_id.lower() != user_id.lower():
                    ts = await get_meeting_transcripts(credential, rec_organizer_id, online_meeting_id)
                meeting_transcripts_cache[online_meeting_id] = ts
            meeting_transcripts = meeting_transcripts_cache[online_meeting_id]

        if meeting_transcripts:
            matched, match_delta_min = _pick_meeting_transcript(
                meeting_transcripts, rec_start_dt, rec_end_dt, used_transcript_ids
            )

        if not matched and online_meeting_id:
            if bulk_transcripts is None:
                bulk_transcripts = await get_all_transcripts_for_user(credential, user_id)
                if not bulk_transcripts and rec_organizer_id and rec_organizer_id.lower() != user_id.lower():
                    bulk_transcripts = await get_all_transcripts_for_user(credential, rec_organizer_id)
                bulk_transcripts = bulk_transcripts or []
            same_meeting = [
                t for t in bulk_transcripts
                if (t.get("meetingId") or "") == online_meeting_id
            ]
            if same_meeting:
                matched, match_delta_min = _pick_meeting_transcript(
                    same_meeting, rec_start_dt, rec_end_dt, used_transcript_ids
                )

        if matched:
            used_transcript_ids.add(matched.get("id", ""))
            logging.info(
                "Direct time-match: transcript %s (delta=%.0f min)",
                matched.get("id", "")[:40], match_delta_min if match_delta_min is not None else -1,
            )
        else:
            reason = (
                "No transcript found for this meeting"
                if online_meeting_id else
                "No online meeting id resolved for recording"
            )
            results["failed"].append({"recording": rec_name, "reason": reason})
            continue

        try:
            transcript_id = matched.get("id", "")
            meeting_id    = rec.get("online_meeting_id", "") or matched.get("meetingId", "")
            content_url   = matched.get("transcriptContentUrl", "")
            vtt_content   = None

            if not force and not mark_as_processed(user_id, rec_id, rec_name, "in-progress"):
                logging.info("SKIP (claimed by concurrent run): %s", rec_name)
                continue

            if content_url:
                access_token = await get_access_token(credential)
                if access_token:
                    hdrs     = {"Authorization": f"Bearer {access_token}", "Accept": "text/vtt"}
                    url_v1   = content_url.replace("/beta/", "/v1.0/") if "/beta/" in content_url else content_url
                    url_beta = content_url.replace("/v1.0/", "/beta/") if "/v1.0/" in content_url else content_url
                    variants = list(dict.fromkeys([content_url, url_v1, url_beta]))
                    async with httpx.AsyncClient() as client:
                        for variant in variants:
                            resp = await client.get(variant, headers=hdrs, timeout=30.0)
                            if resp.status_code == 200:
                                vtt_content = resp.text
                                break
                            logging.warning("Content URL failed (%s): %s", resp.status_code, variant[:80])

            if not vtt_content and meeting_id and transcript_id:
                vtt_content = await download_transcript_content(credential, user_id, meeting_id, transcript_id)

            if not vtt_content:
                results["failed"].append({"recording": rec_name, "reason": "Transcript download failed"})
                continue

            content_hash = hashlib.sha256(vtt_content.encode("utf-8")).hexdigest()[:16]
            if content_hash in used_content_hashes:
                logging.warning(
                    "SKIP (duplicate transcript content, hash=%s): %s", content_hash, rec_name
                )
                results["failed"].append({
                    "recording": rec_name,
                    "reason": "Duplicate transcript content — already processed in this batch",
                })
                continue
            used_content_hashes.add(content_hash)

            base_name        = build_output_basename(rec_name, meeting_date)
            docx_file        = f"{base_name}.docx"
            transcript_lines = parse_vtt_to_lines(vtt_content)
            if not transcript_lines:
                results["failed"].append({"recording": rec_name, "reason": "VTT parsing returned no lines"})
                continue

            meeting_title = rec_name.replace(".mp4", "").replace(".webm", "").replace(".mkv", "").replace("_", " ").replace("-Meeting Recording", "").strip()

            roster_attendees: List[str] = []
            if ATTENDEE_SOURCE == "attendance":
                roster_attendees = await get_meeting_attendees(
                    credential,
                    user_id,
                    meeting_id,
                    rec_start_dt=rec_start_dt,
                    rec_end_dt=rec_end_dt,
                )
                if (not roster_attendees and rec_organizer_id
                        and rec_organizer_id.lower() != user_id.lower()):
                    roster_attendees = await get_meeting_attendees(
                        credential,
                        rec_organizer_id,
                        meeting_id,
                        rec_start_dt=rec_start_dt,
                        rec_end_dt=rec_end_dt,
                    )
            elif ATTENDEE_SOURCE == "calendar":
                roster_attendees = await get_calendar_event_attendees(
                    credential,
                    user_id,
                    ical_uid=rec.get("ical_uid"),
                    start_dt=rec.get("start_dt"),
                )

            _final_for_log = _resolve_final_attendees(roster_attendees, transcript_lines)
            if roster_attendees:
                _att_source = ATTENDEE_SOURCE
            else:
                _att_source = (
                    "speakers" if ATTENDEE_SOURCE == "speakers"
                    else f"{ATTENDEE_SOURCE}->empty, fell back to speakers"
                )
            logging.info(
                "ATTENDEES for '%s' -> %d resolved [source=%s]: %s",
                rec_name,
                len(_final_for_log),
                _att_source,
                ", ".join(_final_for_log) if _final_for_log else "(none)",
            )
            logging.info("OUTPUT FILENAME would be: %s", docx_file)

            sp_path = None
            if USE_SHAREPOINT:
                try:
                    docx_bytes = build_docx_bytes(
                        meeting_title,
                        meeting_date,
                        transcript_lines,
                        roster_attendees=roster_attendees,
                    )
                    sp_path    = await upload_to_sharepoint(credential, docx_file, docx_bytes, meeting_date)
                except Exception as docx_err:
                    logging.error("DOCX build/upload failed for %s: %s", rec_name, docx_err)
            else:
                final_attendees = _resolve_final_attendees(roster_attendees, transcript_lines)
                preview = "\n".join(f"{e['speaker']}: {e['text']}" for e in transcript_lines[:10])
                logging.info(
                    "SharePoint disabled (DRY RUN).\n"
                    "  Output file : %s\n"
                    "  Attendees   : %s\n"
                    "  Transcript preview (first 10 lines):\n%s",
                    docx_file,
                    ", ".join(final_attendees) if final_attendees else "(none)",
                    preview,
                )

            if STORAGE_CONN_STR:
                try:
                    tc = _get_table_client()
                    if tc:
                        tc.upsert_entity({
                            "PartitionKey":   user_id,
                            "RowKey":         rec_id,
                            "ProcessedAt":    datetime.now(timezone.utc).isoformat(),
                            "RecordingName":  rec_name,
                            "SharePointPath": sp_path or "no-sharepoint",
                        })
                except Exception as upd_err:
                    logging.warning("Failed to update sp_path in table: %s", upd_err)

            results["successful"].append({"recording": rec_name, "transcript_id": transcript_id, "sp_path": sp_path})
        except Exception as exc:
            logging.exception("Error processing '%s': %s", rec_name, exc)
            results["failed"].append({"recording": rec_name, "reason": str(exc)})

    return results


async def get_all_users(graph_client: GraphServiceClient) -> List[str]:
    user_ids: List[str] = []
    page = await graph_client.users.get()
    while page:
        if page.value:
            for user in page.value:
                if getattr(user, "id", None):
                    user_ids.append(user.id)
        next_link = getattr(page, "odata_next_link", None)
        if not next_link:
            break
        page = await graph_client.users.with_url(next_link).get()
    return user_ids


async def get_group_members(graph_client: GraphServiceClient, group_id: str) -> List[str]:
    member_ids: List[str] = []
    page = await graph_client.groups.by_group_id(group_id).members.get()
    while page:
        if page.value:
            for m in page.value:
                odata_type = getattr(m, "odata_type", None) or getattr(m, "@odata_type", None)
                if odata_type == "#microsoft.graph.user" and getattr(m, "id", None):
                    member_ids.append(m.id)
                elif getattr(m, "id", None) and not odata_type:
                    member_ids.append(m.id)
        next_link = getattr(page, "odata_next_link", None)
        if not next_link:
            break
        page = await graph_client.groups.by_group_id(group_id).members.with_url(next_link).get()
    return member_ids


async def get_target_user_ids(
    graph_client: GraphServiceClient,
    mode: UserMode,
    single_user_id: Optional[str] = None,
    group_ids: Optional[List[str]] = None,
) -> List[str]:
    if mode == "single":
        if not single_user_id:
            raise ValueError("single_user_id required")
        return [single_user_id]
    if mode == "all":
        return await get_all_users(graph_client)
    if mode == "users":
        if not group_ids:
            raise ValueError("RECORDING_GROUP_IDS required")
        return group_ids
    if mode == "groups":
        if not group_ids:
            raise ValueError("group_ids required")
        user_ids: set = set()
        for gid in group_ids:
            members = await get_group_members(graph_client, gid)
            user_ids.update(members)
        return list(user_ids)
    raise ValueError(f"Unknown mode: {mode}")


async def run_job():
    logging.info("=== Starting transcript extraction job ===")
    logging.info("Allowed template IDs: %s", ALLOWED_TEMPLATE_IDS)

    missing = [k for k, v in [("TENANT_ID", TENANT_ID), ("CLIENT_ID", CLIENT_ID), ("CLIENT_SECRET", CLIENT_SECRET)] if not v]
    if missing:
        raise EnvironmentError(f"Missing required env var(s): {', '.join(missing)}")

    credential   = ClientSecretCredential(tenant_id=TENANT_ID, client_id=CLIENT_ID, client_secret=CLIENT_SECRET)
    graph_client = GraphServiceClient(credentials=credential, scopes=["https://graph.microsoft.com/.default"])

    mode, single_user_id, group_ids = get_user_mode_from_env()
    target_user_ids = await get_target_user_ids(graph_client, mode=mode, single_user_id=single_user_id, group_ids=group_ids)
    logging.info("Processing %d user(s), mode=%s", len(target_user_ids), mode)

    overall = {"successful": [], "failed": []}
    used_content_hashes: set = set()

    for user_id in target_user_ids:
        logging.info("--- Processing user %s ---", user_id)
        try:
            user  = await graph_client.users.by_user_id(user_id).get()
            name  = getattr(user, "display_name", "Unknown")
            email = getattr(user, "mail", None) or getattr(user, "user_principal_name", "")
            logging.info("User: %s | %s | %s", user_id, name, email)
        except Exception as exc:
            logging.error("Skipping user %s: %s", user_id, exc)
            overall["failed"].append({"recording": f"user:{user_id}", "reason": str(exc)})
            continue

        try:
            recordings, _ = await get_user_recordings(graph_client, credential, user_id)
            if not recordings:
                continue
            results = await process_recordings(graph_client, credential, user_id, recordings, used_content_hashes)
            overall["successful"].extend(results["successful"])
            overall["failed"].extend(results["failed"])
        except Exception as exc:
            logging.exception("Error processing user %s: %s", user_id, exc)
            overall["failed"].append({"recording": f"user:{user_id}", "reason": str(exc)})

    logging.info("=== Job complete. PASS=%d  FAIL=%d ===", len(overall["successful"]), len(overall["failed"]))
    if overall["failed"]:
        reason_counts: Dict[str, int] = {}
        for f in overall["failed"]:
            r = f.get("reason", "unknown")
            reason_counts[r] = reason_counts.get(r, 0) + 1
        logging.warning(
            "Failure reasons: %s",
            "; ".join(f"{r} x{n}" for r, n in sorted(reason_counts.items(), key=lambda kv: -kv[1])),
        )

    return overall


def _setup_logging() -> None:
    root = logging.getLogger()
    if not root.handlers:
        logging.basicConfig(level=logging.INFO, format="%(asctime)s - %(levelname)s - %(message)s", datefmt="%Y-%m-%d %H:%M:%S")
    else:
        root.setLevel(logging.INFO)
    for noisy in ("azure", "msal", "httpx", "httpcore", "msgraph", "urllib3",
                  "azure.core.pipeline.policies.http_logging_policy"):
        logging.getLogger(noisy).setLevel(logging.WARNING)


def main(mytimer=None) -> dict:
    _setup_logging()
    overall = asyncio.run(run_job())
    print(f"PASS={len(overall['successful'])}  FAIL={len(overall['failed'])}")
    return overall


if __name__ == "__main__":
    _setup_logging()
    overall = asyncio.run(run_job())
    print(f"PASS={len(overall['successful'])}  FAIL={len(overall['failed'])}")